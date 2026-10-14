"""
================================================================================
LIVE DESKTOP CHROME AUTOMATION - COMPANY DETAILS ENRICHMENT
(COMPANY WEBSITE, COMPANY PHONE NUMBER & NO OF EMPLOYEES)
================================================================================

DESCRIPTION:
    Automates Google Chrome directly on your desktop screen so you can watch it
    work live! For each company in your Excel file (e.g. companies_sample.xlsx):
      1. Reads the company name from Excel row-by-row.
      2. Adds 3 new columns if not already present:
            - Column 1: 'Company Website'
            - Column 2: 'Company Phone Number'
            - Column 3: 'No of Employees'
      3. Launches a real, visible Google Chrome window maximized on your desktop.
      4. Visibly types the query into Chrome and searches live:
            - Finds the company's official Website URL.
            - Finds the company's verified Phone Number / Contact Number.
            - Finds the company's Number of Employees (company size / headcount).
      5. Writes the extracted details into the Excel spreadsheet in the respective columns.
      6. Saves progress incrementally every N rows (default: 5).
      7. Safe Resume: Skips rows where all 3 fields are already populated.
      8. Lock Protection: If the Excel file is open in Microsoft Excel, saves to
         a backup file (e.g. companies_sample_updated.xlsx) to prevent data loss.

HOW TO RUN:
    python enrich_company_details.py

OPTIONS:
    Test on just 5 companies first:
        python enrich_company_details.py --limit 5

    Specify a different Excel file:
        python enrich_company_details.py --file "C:/path/to/my_file.xlsx"

    Run in background without opening visible Chrome window:
        python enrich_company_details.py --headless

    Custom auto-save interval and delay:
        python enrich_company_details.py --save-interval 5 --delay 1.5

DEPENDENCIES:
    pip install openpyxl selenium webdriver-manager requests python-dotenv
================================================================================
"""

import os
import re
import sys
import time
import random
import logging
import argparse
from typing import Dict, Optional, Tuple, List, Any
from urllib.parse import urlparse, quote_plus, parse_qs

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.styles import Font, PatternFill, Alignment
from dotenv import load_dotenv

from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys

load_dotenv()

# Logging configuration
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    datefmt="%H:%M:%S"
)
logger = logging.getLogger("company_details_enricher")

# Silence third-party loggers
for mod in ("selenium", "urllib3", "webdriver_manager"):
    logging.getLogger(mod).setLevel(logging.WARNING)


# Header matching candidates for fuzzy column detection
COMPANY_HEADER_CANDIDATES = [
    "name", "company", "company name", "company_name", "organization",
    "firm", "account name", "companies"
]

WEBSITE_HEADER_CANDIDATES = [
    "company website", "companys website", "company's website",
    "website", "company url", "web address", "official website", "url"
]

PHONE_HEADER_CANDIDATES = [
    "company phone number", "comapny phone number", "company phone",
    "phone number", "phone", "contact number", "telephone", "mobile", "tel"
]

EMPLOYEES_HEADER_CANDIDATES = [
    "no of employees", "number of employees", "employees", "employee count",
    "company size", "headcount", "no. of employees", "total employees"
]

# Domains to ignore when looking for official company website
IGNORE_DOMAINS = {
    'linkedin.com', 'facebook.com', 'twitter.com', 'x.com', 'instagram.com', 'youtube.com',
    'wikipedia.org', 'zaubacorp.com', 'tofler.in', 'instafinancials.com', 'justdial.com',
    'indiamart.com', 'glassdoor.com', 'ambitionbox.com', 'indeed.com', 'crunchbase.com',
    'zoominfo.com', 'yellowpages.com', 'quora.com', 'naukri.com', 'mapsofindia.com',
    'google.com', 'duckduckgo.com', 'bing.com', 'bloomberg.com', 'tracxn.com',
    'tutorialspoint.com', 'geeksforgeeks.org', 'w3schools.com', 'javatpoint.com',
    'reddit.com', 'pinterest.com', 'github.com', 'medium.com', 'tradeindia.com',
    'quickr.com', 'sulekha.com', 'mouthshut.com', 'economictimes.indiatimes.com',
    'moneycontrol.com', 'livemint.com', 'business-standard.com', 'thehindubusinessline.com',
    'pitchbook.com', 'owler.com', 'apollo.io', 'signalhire.com', 'rocketreach.co',
    'yahoo.com', 'msn.com', 'apple.com', 'microsoft.com'
}


def clean_company_name(name: Any) -> str:
    """Cleans whitespace, quotes, and punctuation from company names."""
    if not name:
        return ""
    cleaned = str(name).strip().strip('\'"“”')
    cleaned = re.sub(r'\s+', ' ', cleaned)
    return cleaned


def extract_brand_name(company_name: str) -> str:
    """Extracts clean brand name without corporate legal suffixes (Pvt Ltd, LLC, etc.)."""
    clean = clean_company_name(company_name)
    suffix_free = re.sub(
        r'\b(?:pvt\.?\s*ltd\.?|private\s+limited|p\s+ltd|ltd\.?|limited|inc\.?|llc|corp\.?|co\.?)\b',
        '', clean, flags=re.IGNORECASE
    ).strip()
    suffix_free = re.sub(r'[\s.,&;/-]+$', '', suffix_free).strip()
    return suffix_free if len(suffix_free) >= 3 else clean


def clean_phone_number(raw_phone: str) -> Optional[str]:
    """Cleans and standardizes phone number formats."""
    if not raw_phone:
        return None
    phone = raw_phone.strip()
    digits = re.sub(r'\D', '', phone)

    if not (8 <= len(digits) <= 15):
        return None

    # Toll-Free numbers
    if digits.startswith("1800") and len(digits) in (10, 11):
        return f"1800 {digits[4:7]} {digits[7:]}"
    if digits.startswith("800") and len(digits) == 11:
        return f"1800 {digits[3:6]} {digits[6:]}"

    # Indian Mobile (10 digits starting with 6, 7, 8, 9)
    if len(digits) == 10 and digits[0] in "6789":
        return f"+91 {digits[:5]} {digits[5:]}"
    if len(digits) == 12 and digits.startswith("91") and digits[2] in "6789":
        return f"+91 {digits[2:7]} {digits[7:]}"

    # Landline with STD code 044 (Chennai), 080 (Bangalore), 022 (Mumbai), 011 (Delhi)
    for std in ("044", "080", "022", "011", "020", "040", "033", "0124", "0120"):
        if digits.startswith(std):
            rest = digits[len(std):]
            if len(rest) == 8:
                return f"{std}-{rest[:4]} {rest[4:]}"
        if digits.startswith("91" + std[1:]):
            rest = digits[len("91" + std[1:]):]
            if len(rest) == 8:
                return f"+91 {std[1:]} {rest[:4]} {rest[4:]}"

    # General standard formatting if starts with +
    if phone.startswith("+"):
        return phone
    return phone


def unwrap_google_url(url: str) -> str:
    """Unwraps Google redirect URLs (e.g. https://www.google.com/url?q=...) to real target URL."""
    if not url:
        return ""
    if "google.com/url" in url:
        try:
            parsed = urlparse(url)
            qs = parse_qs(parsed.query)
            if 'q' in qs and qs['q']:
                return qs['q'][0]
            if 'url' in qs and qs['url']:
                return qs['url'][0]
        except Exception:
            pass
    return url


def extract_website_from_urls(brand: str, urls: List[str]) -> Optional[str]:
    """
    Selects the best official company website from a list of candidate URLs.
    Scores URLs higher if domain contains the brand tokens.
    Normalizes internal portal subdomains (hr., login., smartpay., etc.) to root domain.
    """
    if not urls:
        return None

    brand_tokens = [t.lower() for t in re.split(r'[^a-zA-Z0-9]', brand) if len(t) >= 3]
    best_url = None
    best_score = -1

    portal_subdomains = (
        "smartpay.", "smarthro.", "hr.", "login.", "portal.", "mail.",
        "webmail.", "beta.", "app.", "staging.", "auth.", "sso."
    )

    for raw_url in urls:
        if not raw_url:
            continue
        unwrapped = unwrap_google_url(raw_url.strip())
        if not unwrapped.startswith("http"):
            continue
        try:
            parsed = urlparse(unwrapped)
            netloc = parsed.netloc.lower()
            if netloc.startswith("www."):
                netloc = netloc[4:]

            # Exclude known directories and search portals
            if any(ign in netloc for ign in IGNORE_DOMAINS):
                continue

            # Exclude login, auth, or signup paths
            path = parsed.path.lower()
            if any(p in path for p in ("/login", "/signin", "/signup", "/register", "/cart", "/search")):
                continue

            # Clean portal subdomains to main website domain
            cleaned_netloc = netloc
            for p_sub in portal_subdomains:
                if cleaned_netloc.startswith(p_sub):
                    cleaned_netloc = cleaned_netloc[len(p_sub):]
                    break

            root_url = f"{parsed.scheme}://www.{cleaned_netloc}" if not cleaned_netloc.startswith("www.") else f"{parsed.scheme}://{cleaned_netloc}"

            # Calculate match score based on brand keywords
            score = 1
            for token in brand_tokens:
                if token in cleaned_netloc:
                    score += 20
                elif token in path:
                    score += 3

            if score > best_score:
                best_score = score
                best_url = root_url

        except Exception:
            continue

    return best_url


def extract_phone_from_text(text: str) -> Optional[str]:
    """Extracts phone number (Indian STD landline, Toll-Free, or Mobile) from text."""
    if not text:
        return None

    # 1. Labeled phone numbers (e.g. Phone: +91 44 2827 8463, Tel: 044-24520125, etc.)
    m_labeled = re.search(
        r'(?:tel|telephone|phone|ph|call|mobile|mob|contact|toll\s*free|helpdesk|customer\s*care)[\s.:#-]*(\+?[0-9]{1,4}[\s-]?(?:\(?[0-9]{2,5}\)?[\s-]?)?[0-9]{3,5}[\s-]?[0-9]{3,5})',
        text, re.IGNORECASE
    )
    if m_labeled:
        cand = clean_phone_number(m_labeled.group(1))
        if cand:
            return cand

    # 2. Toll-Free numbers (e.g. 1800 200 3000, 1860-XXX-XXXX)
    m_tf = re.search(r'\b(1800[\s-]?[0-9]{3}[\s-]?[0-9]{3,4}|1860[\s-]?[0-9]{3}[\s-]?[0-9]{3,4})\b', text)
    if m_tf:
        cand = clean_phone_number(m_tf.group(1))
        if cand:
            return cand

    # 3. Standard Indian Landline with STD code (e.g. 044-XXXXXXXX, 080-XXXXXXXX, +91 44 XXXXXXXX)
    m_std = re.search(r'\b(0[1-9][0-9]{1,3}[-\s]?[2-9][0-9]{5,7}|\+91[\s-]?[0-9]{2,4}[\s-]?[0-9]{6,8})\b', text)
    if m_std:
        cand = clean_phone_number(m_std.group(1))
        if cand:
            return cand

    # 4. Standard 10-digit mobile number with +91 or 0
    m_mob = re.search(r'\b(?:\+91[\s-]?)?[6-9][0-9]{4}[\s-]?[0-9]{5}\b', text)
    if m_mob:
        cand = clean_phone_number(m_mob.group(0))
        if cand:
            return cand

    return None


def extract_employees_from_text(text: str) -> Optional[str]:
    """Extracts number of employees / company size from snippet text."""
    if not text:
        return None

    # 1. Standard employee range: e.g. '51-200 employees', '1,001-5,000 employees', '10,001+ employees'
    m_range = re.search(
        r'\b([0-9]{1,3}(?:,[0-9]{3})*\s*-\s*[0-9]{1,3}(?:,[0-9]{3})*|\b[0-9]{1,3}(?:,[0-9]{3})*\+)\s*(?:employees|employee|staff|people|workforce)\b',
        text, re.IGNORECASE
    )
    if m_range:
        return f"{m_range.group(1).strip()} employees"

    # 2. Sentences like 'employee count at COMPANY is 812' or 'headcount of 812' or 'team size is 450'
    m_count_sentence = re.search(
        r'(?:employee count|headcount|workforce|team size)(?:[^\.\n]{0,80}?(?:is|was|of|:|=))\s*([0-9]{1,3}(?:,[0-9]{3})*|[0-9]+(?:\s*-\s*[0-9]+|\+)?)\b',
        text, re.IGNORECASE
    )
    if m_count_sentence:
        val = m_count_sentence.group(1).strip()
        return f"{val} employees" if "employee" not in val.lower() else val

    # 3. Labeled company size: 'Company size: 501-1,000 employees' or 'Headcount: 250 employees'
    m_label = re.search(
        r'(?:company size|headcount|workforce|employee count|team size|size)[:\s]+([0-9,+-]+(?:\s*(?:to|-)\s*[0-9,+-]+)?(?:\s*employees)?)',
        text, re.IGNORECASE
    )
    if m_label:
        val = m_label.group(1).strip()
        if 'employee' not in val.lower():
            val += " employees"
        return val

    # 4. Direct count: e.g. '4,820 employees', '300 employees'
    m_count = re.search(r'\b([0-9]{1,3}(?:,[0-9]{3})*)\s+(?:employees|employee)\b', text, re.IGNORECASE)
    if m_count:
        val = m_count.group(1).strip()
        try:
            digits = int(val.replace(',', ''))
            if digits > 0:
                return f"{val} employees"
        except ValueError:
            pass

    return None


class LiveBrowserAutomator:
    """Automates visible Google Chrome directly on user's desktop using Google Search ONLY."""

    def __init__(self, headless: bool = False):
        self.headless = headless
        self.driver: Optional[webdriver.Chrome] = None
        self.profile_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), "chrome_google_profile"))

    def start(self) -> None:
        """Launches Google Chrome maximized on desktop using Google Search."""
        logger.info("Launching Google Chrome browser on desktop (Google Search)...")
        os.makedirs(self.profile_dir, exist_ok=True)

        options = Options()
        if self.headless:
            options.add_argument('--headless=new')
        else:
            options.add_argument('--start-maximized')

        options.add_argument(f'--user-data-dir={self.profile_dir}')
        options.add_argument('--disable-blink-features=AutomationControlled')
        options.add_experimental_option('excludeSwitches', ['enable-automation'])
        options.add_experimental_option('useAutomationExtension', False)
        options.add_argument(
            "user-agent=Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
            "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/130.0.0.0 Safari/537.36"
        )

        self.driver = webdriver.Chrome(options=options)
        try:
            self.driver.execute_cdp_cmd('Page.addScriptToEvaluateOnNewDocument', {
                'source': '''
                    Object.defineProperty(navigator, 'webdriver', { get: () => undefined });
                    window.navigator.chrome = { runtime: {} };
                    Object.defineProperty(navigator, 'plugins', { get: () => [1, 2, 3, 4, 5] });
                    Object.defineProperty(navigator, 'languages', { get: () => ['en-US', 'en'] });
                '''
            })
        except Exception:
            pass

        self.driver.get('https://www.google.com')
        time.sleep(1.5)
        self.dismiss_consent()

    def dismiss_consent(self) -> None:
        """Dismisses Google cookie/consent dialog if present."""
        for selector in ['button#L2AGLb', 'button#W0wltc', 'button[aria-label*="Accept"]', 'form[action*="consent"] button']:
            try:
                btn = self.driver.find_element(By.CSS_SELECTOR, selector)
                btn.click()
                time.sleep(1.0)
                break
            except Exception:
                pass

    def check_and_handle_captcha(self) -> None:
        """
        Detects if Google presents a reCAPTCHA / verification screen.
        Alerts user and pauses until user solves it on screen.
        """
        if not self.driver:
            return

        if 'sorry' in self.driver.current_url or "detected unusual traffic" in self.driver.page_source.lower():
            logger.warning("\n" + "!" * 70)
            logger.warning("   [GOOGLE VERIFICATION DETECTED ON YOUR CHROME SCREEN]")
            logger.warning("   Google has displayed a verification challenge.")
            logger.warning("   Please click the 'I'm not a robot' checkbox in the open Chrome window.")
            logger.warning("   The script will automatically detect completion and resume immediately!")
            logger.warning("!" * 70 + "\n")

            while 'sorry' in self.driver.current_url or "detected unusual traffic" in self.driver.page_source.lower():
                time.sleep(2)

            logger.info("   -> Verification completed! Continuing Google search...")
            time.sleep(2.0)

    def close(self) -> None:
        """Quits the browser."""
        if self.driver:
            try:
                self.driver.quit()
            except Exception:
                pass
            self.driver = None

    def type_and_search(self, query: str) -> None:
        """Visibly types query into Google search box with typing simulation."""
        self.check_and_handle_captcha()

        box = None
        for selector in ['textarea[name="q"]', 'input[name="q"]', '[name="q"]']:
            try:
                box = self.driver.find_element(By.CSS_SELECTOR, selector)
                break
            except Exception:
                continue

        if not box:
            self.driver.get('https://www.google.com')
            time.sleep(1.5)
            self.dismiss_consent()
            for selector in ['textarea[name="q"]', 'input[name="q"]', '[name="q"]']:
                try:
                    box = self.driver.find_element(By.CSS_SELECTOR, selector)
                    break
                except Exception:
                    continue

        if box:
            box.send_keys(Keys.CONTROL + "a")
            box.send_keys(Keys.BACKSPACE)
            time.sleep(0.1)

            # Visibly type character by character
            for char in query:
                box.send_keys(char)
                time.sleep(0.012)

            box.send_keys(Keys.RETURN)
            time.sleep(2.5)
        else:
            # Fallback direct Google search URL
            self.driver.get(f"https://www.google.com/search?q={quote_plus(query)}&hl=en")
            time.sleep(2.5)

        self.check_and_handle_captcha()

    def search_website(self, company_name: str) -> Optional[str]:
        """
        Visibly searches for the company's official website using Google Search only.
        Extracts from Knowledge Graph or organic Google search results.
        """
        if not self.driver:
            self.start()

        clean_comp = clean_company_name(company_name)
        brand = extract_brand_name(clean_comp)

        # Primary search: type the company name directly in Google Search: [company name]
        queries = [clean_comp]
        if brand.lower() != clean_comp.lower():
            queries.append(brand)
        queries.append(f"{brand} official website")

        for query in queries:
            logger.info(f"   [Google Chrome] Searching Website: {query}")
            try:
                self.type_and_search(query)

                candidate_urls = []

                # 1. Knowledge Panel Website button (highest confidence)
                kp_buttons = self.driver.find_elements(
                    By.CSS_SELECTOR,
                    'a[aria-label*="Website"], a[role="button"][href*="http"], a[data-ved]'
                )
                for btn in kp_buttons:
                    txt = btn.text.strip().lower()
                    aria = (btn.get_attribute("aria-label") or "").lower()
                    if "website" in txt or "website" in aria:
                        raw_href = btn.get_attribute("href")
                        unwrapped = unwrap_google_url(raw_href)
                        if unwrapped and "google.com" not in unwrapped:
                            candidate_urls.insert(0, unwrapped)

                # 2. Organic Google search result links
                organic_links = self.driver.find_elements(
                    By.CSS_SELECTOR,
                    'div#search a[href], div#rso a[href], div.g a[href], a:has(h3)'
                )
                for a in organic_links[:15]:
                    raw_href = a.get_attribute("href")
                    unwrapped = unwrap_google_url(raw_href)
                    if unwrapped and "google.com" not in unwrapped and not unwrapped.startswith("javascript:"):
                        candidate_urls.append(unwrapped)

                website = extract_website_from_urls(brand, candidate_urls)
                if website:
                    return website

            except Exception as e:
                logger.warning(f"Google website search error for {query}: {e}")

        return None

    def search_phone(self, company_name: str) -> Optional[str]:
        """
        Visibly searches for the company's contact phone number using Google Search only.
        Extracts from Google Knowledge Panel / Business listing or organic snippets.
        """
        if not self.driver:
            self.start()

        clean_comp = clean_company_name(company_name)
        brand = extract_brand_name(clean_comp)

        queries = [
            f'{brand} Chennai phone number contact',
            f'{clean_comp} contact phone number customer care',
            f'{brand} phone number office'
        ]

        for query in queries:
            logger.info(f"   [Google Chrome] Searching Phone: {query}")
            try:
                self.type_and_search(query)

                # 1. Check for tel: links first
                tel_elements = self.driver.find_elements(By.CSS_SELECTOR, 'a[href^="tel:"]')
                for tel_el in tel_elements:
                    href = tel_el.get_attribute('href') or ''
                    tel_num = href.replace('tel:', '').strip()
                    cleaned_p = clean_phone_number(tel_num)
                    if cleaned_p:
                        return cleaned_p

                # 2. Check Google Knowledge Panel phone blocks
                kp_phone_elements = self.driver.find_elements(
                    By.CSS_SELECTOR,
                    '[data-attrid*="phone"], [aria-label*="Call"], [data-dtype="d3ph"], div#rhs'
                )
                for el in kp_phone_elements:
                    txt = el.text.strip()
                    phone = extract_phone_from_text(txt)
                    if phone:
                        return phone

                # 3. Scan organic Google search results & snippets
                snippets = self.driver.find_elements(
                    By.CSS_SELECTOR,
                    'div.VwiC3b, div.g, div[data-hveid], div#rso'
                )
                for s in snippets[:8]:
                    txt = s.text.strip()
                    phone = extract_phone_from_text(txt)
                    if phone:
                        return phone

            except Exception as e:
                logger.warning(f"Google phone search error for {query}: {e}")

        return None

    def search_employees(self, company_name: str) -> Optional[str]:
        """
        Visibly searches for the company's employee count (headcount / size) using Google Search only.
        Extracts from Google Knowledge Panel or organic Google snippets.
        """
        if not self.driver:
            self.start()

        clean_comp = clean_company_name(company_name)
        brand = extract_brand_name(clean_comp)

        # Required search format: [company name] no of employees
        queries = [f"{clean_comp} no of employees"]
        if brand.lower() != clean_comp.lower():
            queries.append(f"{brand} no of employees")

        for query in queries:
            logger.info(f"   [Google Chrome] Searching Employees: {query}")
            try:
                self.type_and_search(query)

                # 1. Check Google Knowledge Panel / Featured Answer Box / Attribute row
                kp_elements = self.driver.find_elements(
                    By.CSS_SELECTOR,
                    '[data-attrid*="employee"], [data-attrid*="size"], div.kp-header, div.Z0LcW, div#rhs'
                )
                for el in kp_elements:
                    txt = el.text.strip()
                    employees = extract_employees_from_text(txt)
                    if employees:
                        return employees

                # 2. Scan Google organic search result snippets (LinkedIn, AmbitionBox, Glassdoor, etc.)
                snippets = self.driver.find_elements(
                    By.CSS_SELECTOR,
                    'div.VwiC3b, div.g, div[data-hveid], div#rso'
                )
                for s in snippets[:10]:
                    txt = s.text.strip()
                    employees = extract_employees_from_text(txt)
                    if employees:
                        return employees

            except Exception as e:
                logger.warning(f"Google employees search error for {query}: {e}")

        return None


class CompanyDetailsEnricher:
    """Manages Excel file enrichment with live Chrome visualization for Website, Phone & Employees."""

    def __init__(
        self,
        file_path: str,
        sheet_name: Optional[str] = None,
        headless: bool = False,
        save_interval: int = 5,
        delay_seconds: float = 1.5,
        limit: Optional[int] = None
    ):
        self.file_path = os.path.abspath(file_path)
        self.sheet_name = sheet_name
        self.headless = headless
        self.save_interval = max(1, save_interval)
        self.delay_seconds = delay_seconds
        self.limit = limit

        self.wb: Optional[openpyxl.Workbook] = None
        self.sheet: Optional[Worksheet] = None
        self.col_map: Dict[str, int] = {}
        self.is_dirty: bool = False
        self.automator = LiveBrowserAutomator(headless=self.headless)

        self.stats = {
            "total_rows": 0,
            "empty_rows": 0,
            "skipped_already_done": 0,
            "processed": 0,
            "website_found": 0,
            "phone_found": 0,
            "employees_found": 0,
        }

    def prepare_excel(self) -> None:
        """Loads spreadsheet, maps columns, and adds new columns if needed."""
        if not os.path.exists(self.file_path):
            raise FileNotFoundError(f"Excel file not found at: {self.file_path}")

        logger.info(f"Opening Excel file: {self.file_path}")
        self.wb = openpyxl.load_workbook(self.file_path)

        if self.sheet_name and self.sheet_name in self.wb.sheetnames:
            self.sheet = self.wb[self.sheet_name]
        else:
            self.sheet = self.wb.active

        logger.info(f"Using Sheet: '{self.sheet.title}' (Total Rows: {self.sheet.max_row})")

        headers: Dict[int, str] = {}
        for c in range(1, self.sheet.max_column + 1):
            val = self.sheet.cell(1, c).value
            if val is not None:
                headers[c] = str(val).strip().lower()

        def match_col(candidates: List[str]) -> Optional[int]:
            for c, text in headers.items():
                for cand in candidates:
                    if cand == text or text.startswith(cand):
                        return c
            return None

        # 1. Company Name Column
        company_col = match_col(COMPANY_HEADER_CANDIDATES) or 2  # default to col 2 if col 1 is S.NO
        self.col_map["company"] = company_col

        # 2. Target Columns: Website, Phone Number, No of Employees
        targets = [
            ("website", WEBSITE_HEADER_CANDIDATES, "Company Website", 35),
            ("phone", PHONE_HEADER_CANDIDATES, "Company Phone Number", 26),
            ("employees", EMPLOYEES_HEADER_CANDIDATES, "No of Employees", 24),
        ]

        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        header_align = Alignment(horizontal="center", vertical="center")

        curr_max = self.sheet.max_column
        for key, candidates, title, col_width in targets:
            c = match_col(candidates)
            if c:
                self.col_map[key] = c
                logger.info(f"Found existing column '{title}' at Column {openpyxl.utils.get_column_letter(c)}")
            else:
                curr_max += 1
                new_cell = self.sheet.cell(1, curr_max, value=title)
                new_cell.font = header_font
                new_cell.fill = header_fill
                new_cell.alignment = header_align
                col_letter = openpyxl.utils.get_column_letter(curr_max)
                self.sheet.column_dimensions[col_letter].width = col_width
                self.col_map[key] = curr_max
                headers[curr_max] = title.lower()
                self.is_dirty = True
                logger.info(f"Appended new column '{title}' at Column {col_letter}")

        logger.info(
            "Columns mapped: Company=%s, Website=%s, Phone=%s, Employees=%s",
            openpyxl.utils.get_column_letter(self.col_map["company"]),
            openpyxl.utils.get_column_letter(self.col_map["website"]),
            openpyxl.utils.get_column_letter(self.col_map["phone"]),
            openpyxl.utils.get_column_letter(self.col_map["employees"]),
        )

        if self.is_dirty:
            self.save()

    def process_row(self, row_idx: int) -> bool:
        """Processes a single company row via visible Google Chrome."""
        raw_company = self.sheet.cell(row=row_idx, column=self.col_map["company"]).value
        cleaned_company = clean_company_name(raw_company)

        if not cleaned_company:
            self.stats["empty_rows"] += 1
            return False

        c_website = self.sheet.cell(row=row_idx, column=self.col_map["website"]).value
        c_phone = self.sheet.cell(row=row_idx, column=self.col_map["phone"]).value
        c_employees = self.sheet.cell(row=row_idx, column=self.col_map["employees"]).value

        has_website = bool(c_website and str(c_website).strip())
        has_phone = bool(c_phone and str(c_phone).strip())
        has_employees = bool(c_employees and str(c_employees).strip())

        # Safe Resume: Skip rows already populated with all 3 fields
        if has_website and has_phone and has_employees:
            self.stats["skipped_already_done"] += 1
            return False

        total_rows = self.sheet.max_row
        print("-" * 70)
        logger.info(f"[Row {row_idx}/{total_rows}] Processing Company: '{cleaned_company}'")
        self.stats["processed"] += 1

        fields_changed = False

        # 1. Search Company Website
        if not has_website:
            website = self.automator.search_website(cleaned_company)
            if website:
                self.sheet.cell(row=row_idx, column=self.col_map["website"], value=website)
                logger.info(f"   [FOUND WEBSITE]   -> {website}")
                fields_changed = True
                self.stats["website_found"] += 1
            else:
                logger.info("   [WEBSITE]         -> [Not Found]")
        else:
            w_str = str(c_website).strip()
            # If website has portal subdomain, normalize it
            brand_token = extract_brand_name(cleaned_company)
            normalized_web = extract_website_from_urls(brand_token, [w_str])
            if normalized_web and normalized_web != w_str:
                self.sheet.cell(row=row_idx, column=self.col_map["website"], value=normalized_web)
                logger.info(f"   [NORMALIZED WEBSITE] -> {normalized_web}")
                fields_changed = True
            else:
                logger.info(f"   -> Website already present: {w_str}")

        time.sleep(0.5)

        # 2. Search Company Phone Number
        if not has_phone:
            phone = self.automator.search_phone(cleaned_company)
            if phone:
                self.sheet.cell(row=row_idx, column=self.col_map["phone"], value=phone)
                logger.info(f"   [FOUND PHONE]     -> {phone}")
                fields_changed = True
                self.stats["phone_found"] += 1
            else:
                logger.info("   [PHONE NUMBER]    -> [Not Found]")
        else:
            p_str = str(c_phone).strip()
            cleaned_p = clean_phone_number(p_str)
            if cleaned_p and cleaned_p != p_str:
                self.sheet.cell(row=row_idx, column=self.col_map["phone"], value=cleaned_p)
                logger.info(f"   [FORMATTED PHONE] -> {cleaned_p}")
                fields_changed = True
            else:
                logger.info(f"   -> Phone already present: {p_str}")

        time.sleep(0.5)

        # 3. Search Number of Employees
        if not has_employees:
            employees = self.automator.search_employees(cleaned_company)
            if employees:
                self.sheet.cell(row=row_idx, column=self.col_map["employees"], value=employees)
                logger.info(f"   [FOUND EMPLOYEES] -> {employees}")
                fields_changed = True
                self.stats["employees_found"] += 1
            else:
                logger.info("   [EMPLOYEES]       -> [Not Found]")
        else:
            logger.info(f"   -> Employees already present: {c_employees}")

        if fields_changed:
            self.is_dirty = True

        return True

    def save(self) -> None:
        """Saves Excel file incrementally, with fallback if open in Excel."""
        if self.wb:
            try:
                self.wb.save(self.file_path)
                self.is_dirty = False
                logger.info(f"[EXCEL SAVED] Progress saved to '{os.path.basename(self.file_path)}'")
            except PermissionError:
                base, ext = os.path.splitext(self.file_path)
                fallback_path = f"{base}_updated{ext}"
                logger.warning(
                    f"\n[FILE LOCKED] '{os.path.basename(self.file_path)}' is currently open in Microsoft Excel!\n"
                    f"   -> Saving progress to '{os.path.basename(fallback_path)}' so no data is lost.\n"
                    f"   -> (TIP: Close the file in Excel to allow the script to save directly to the original file.)"
                )
                try:
                    self.wb.save(fallback_path)
                    logger.info(f"[EXCEL SAVED] Progress secured in '{os.path.basename(fallback_path)}'")
                except Exception as e2:
                    logger.error(f"Fallback save error: {e2}")
            except Exception as e:
                logger.error(f"Error saving Excel file: {e}")

    def run(self) -> None:
        """Runs the loop through all Excel rows."""
        self.prepare_excel()
        self.automator.start()

        max_row = self.sheet.max_row
        processed_count = 0
        since_last_save = 0

        logger.info(f"Starting live screen automation for {max_row - 1} rows...")

        try:
            for row_idx in range(2, max_row + 1):
                self.stats["total_rows"] += 1

                if self.limit and processed_count >= self.limit:
                    logger.info(f"Reached limit of {self.limit} companies. Stopping.")
                    break

                try:
                    searched = self.process_row(row_idx)

                    if searched:
                        processed_count += 1
                        since_last_save += 1

                        if since_last_save >= self.save_interval and self.is_dirty:
                            self.save()
                            since_last_save = 0

                        if row_idx < max_row and (not self.limit or processed_count < self.limit):
                            time.sleep(self.delay_seconds)

                except KeyboardInterrupt:
                    logger.warning("\nExecution stopped by user (Ctrl+C). Saving...")
                    break
                except Exception as e:
                    logger.error(f"[Row {row_idx}] Error: {e}")

        finally:
            if self.is_dirty:
                self.save()
            self.automator.close()
            self.print_summary()

    def print_summary(self) -> None:
        """Prints summary statistics."""
        print("\n" + "=" * 65)
        print("          COMPANY DETAILS ENRICHMENT SUMMARY")
        print("=" * 65)
        print(f"  Total Rows Inspected:       {self.stats['total_rows']}")
        print(f"  Empty Rows Skipped:         {self.stats['empty_rows']}")
        print(f"  Already Enriched (Skipped): {self.stats['skipped_already_done']}")
        print(f"  Companies Processed:        {self.stats['processed']}")
        print(f"  Websites Found:             {self.stats['website_found']}")
        print(f"  Phone Numbers Found:        {self.stats['phone_found']}")
        print(f"  Employee Counts Found:      {self.stats['employees_found']}")
        print("=" * 65 + "\n")


def parse_arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Live Desktop Chrome Automation - Company Website, Phone & Employees Enrichment"
    )
    parser.add_argument(
        "--file", "-f",
        default="companies_sample.xlsx",
        help="Path to Excel spreadsheet (default: companies_sample.xlsx)"
    )
    parser.add_argument(
        "--sheet", "-s",
        default=None,
        help="Sheet name to process (default: active sheet)"
    )
    parser.add_argument(
        "--limit", "-l",
        type=int,
        default=None,
        help="Process only first N companies (for testing)"
    )
    parser.add_argument(
        "--delay", "-d",
        type=float,
        default=1.5,
        help="Delay in seconds between searches (default: 1.5)"
    )
    parser.add_argument(
        "--save-interval",
        type=int,
        default=5,
        help="Save Excel workbook every N rows (default: 5)"
    )
    parser.add_argument(
        "--headless",
        action="store_true",
        help="Run Chrome silently in background (headless mode)"
    )
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_arguments()
    excel_path = os.path.abspath(args.file)

    if not os.path.exists(excel_path):
        # Check Desktop fallback
        desktop_candidate = os.path.join(os.path.expanduser("~"), "Desktop", os.path.basename(excel_path))
        if os.path.exists(desktop_candidate):
            excel_path = desktop_candidate
        else:
            print(f"Error: File not found at '{excel_path}'")
            sys.exit(1)

    enricher = CompanyDetailsEnricher(
        file_path=excel_path,
        sheet_name=args.sheet,
        headless=args.headless,
        save_interval=args.save_interval,
        delay_seconds=args.delay,
        limit=args.limit
    )

    enricher.run()
