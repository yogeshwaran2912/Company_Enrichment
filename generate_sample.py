"""
Helper script to generate a sample Excel spreadsheet ('companies_sample.xlsx')
for testing the company enrichment automation script.
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def create_sample_excel(file_path: str = "companies_sample.xlsx"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Companies"

    # Header styling
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    headers = [
        "Company Name",
        "CEO Name",
        "CEO LinkedIn URL",
        "HR Head Name",
        "HR Head LinkedIn URL",
        "Headquarters",
        "Admin / Support Person Name",
        "Admin / Support Person LinkedIn"
    ]

    ws.append(headers)

    # Style header row
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    sample_data = [
        [
            "Microsoft",
            "Satya Nadella",
            "https://www.linkedin.com/in/satyanadella/",
            "Kathleen Hogan",
            "https://www.linkedin.com/in/kathleen-hogan-msft/",
            "Redmond, Washington, USA",
            "",
            ""
        ],
        ["Apple", "", "", "", "", "", "", ""],
        ["NVIDIA", "", "", "", "", "", "", ""],
        ["Stripe", "", "", "", "", "", "", ""],
        ["Airbnb", "", "", "", "", "", "", ""]
    ]

    for row_data in sample_data:
        ws.append(row_data)

    column_widths = {
        "A": 22,
        "B": 24,
        "C": 45,
        "D": 24,
        "E": 45,
        "F": 30,
        "G": 28,
        "H": 45
    }
    for col_letter, width in column_widths.items():
        ws.column_dimensions[col_letter].width = width

    wb.save(file_path)
    print(f"Sample Excel spreadsheet created successfully at: {file_path}")

if __name__ == "__main__":
    create_sample_excel()
