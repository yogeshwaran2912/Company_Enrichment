"""
================================================================================
SHORTEN EXISTING CORE BUSINESS IN EXCEL
================================================================================
Scans an Excel file and shortens long sentence descriptions in the 'Core Business'
column into clean, concise tags (e.g., 'IT Infrastructure', 'Industrial Automation').
================================================================================
"""

import os
import sys
import openpyxl
from core_business_extractor import extract_short_core_business

DEFAULT_FILES = [
    r"C:\Users\ADV PRO\Desktop\company_enrichmentchennai companys.xlsx",
    r"C:\Users\ADV PRO\Desktop\company_enrichmentchennai companys_updated.xlsx"
]

def shorten_excel(file_path: str):
    if not os.path.exists(file_path):
        print(f"[SKIP] File not found: {file_path}")
        return

    print(f"\nProcessing Excel: {file_path}")
    wb = openpyxl.load_workbook(file_path)
    sheet = wb.active

    headers = [sheet.cell(1, c).value for c in range(1, sheet.max_column + 1)]
    comp_col = None
    biz_col = None

    for idx, h in enumerate(headers, start=1):
        if not h:
            continue
        h_str = str(h).strip().lower()
        if h_str in ["name", "company name", "company"]:
            comp_col = idx
        elif h_str in ["core business", "nature of business", "business overview"]:
            biz_col = idx

    if not comp_col or not biz_col:
        print(f"Could not map columns. comp_col={comp_col}, biz_col={biz_col}")
        return

    shortened_count = 0
    for r in range(2, sheet.max_row + 1):
        comp_name = sheet.cell(r, comp_col).value
        biz_val = sheet.cell(r, biz_col).value

        if biz_val and str(biz_val).strip():
            val_str = str(biz_val).strip()
            # If value has more than 5 words or > 40 chars, it's a long sentence
            if len(val_str.split()) > 5 or len(val_str) > 40:
                short = extract_short_core_business(val_str, company_name=str(comp_name or ""))
                if short and short != val_str:
                    print(f"  [Row {r}] '{comp_name}'")
                    print(f"     Old: {val_str[:65]}...")
                    print(f"     New: {short}")
                    sheet.cell(r, biz_col, value=short)
                    shortened_count += 1

    if shortened_count > 0:
        try:
            wb.save(file_path)
            print(f"--> Saved! Shortened {shortened_count} rows in '{os.path.basename(file_path)}'")
        except PermissionError:
            base, ext = os.path.splitext(file_path)
            fallback = f"{base}_shortened{ext}"
            wb.save(fallback)
            print(f"--> Original file was locked. Saved {shortened_count} rows to '{os.path.basename(fallback)}'")
    else:
        print("--> All populated rows are already short and concise!")


if __name__ == "__main__":
    target = sys.argv[1] if len(sys.argv) > 1 else None
    if target:
        shorten_excel(target)
    else:
        for f in DEFAULT_FILES:
            shorten_excel(f)
